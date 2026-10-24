"""P1-28：DOCX/PPTX/PDF 抽图、位置/附件映射、跳过与预算回归。

证据要求（START_S06 §三 P1-28）：

* 三格式各自「非空正文 + 有效图」，断言 ``[图片 N]`` 标记、真实 image record 的
  filename/id/bytes、源文档对应与位置顺序；统计用**实际返回**的 emitted bytes/
  count 与 ``extracted_chars`` 验证，不用 ``try_admit()`` 代替贯通。
* 12/13 每文档、16/17 每轮、1199/1200、4MiB/超 1 字节、12MiB/超 1 字节、
  两个文档共享 turn 额度、重复资产复用且重复位置标记、非支持/矢量图、
  坏 ZIP/坏 XML/外部与穿越关系。
* 不解压到宿主任意路径、不发网络；无图文本与 reading 专用预算行为保留。

本文件只使用临时目录与内存中的最小有效 fixture；生产代码仅在测试证明存在缺陷时
做最小修改（穿越关系见 ``_office_rels`` 的包根约束）。
"""

from __future__ import annotations

import base64
import io
from pathlib import Path
import random
from typing import Any
import zipfile

from docx import Document as DocxDocument
from docx.shared import Inches
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PptxInches
import pymupdf
import pytest

from deeptutor.utils import document_extractor as extractor
from deeptutor.utils.document_images import (
    MAX_IMAGE_BYTES,
    MAX_IMAGES_PER_DOC,
    MAX_TOTAL_IMAGE_BYTES_PER_DOC,
    MIN_IMAGE_BYTES,
    ImageBudget,
    extract_docx_rich,
    extract_pdf_images,
    extract_pptx_rich,
    find_markers,
    reading_image_budget,
)
from tests.phase1 import helpers

_MIME = "image/png"


# ---------------------------------------------------------------------------
# fixtures
# ---------------------------------------------------------------------------
def _png_bytes(
    target: int | None = None, *, side: int = 64, seed: int = 1, color: tuple[int, int, int] | None = None
) -> bytes:
    """A valid raster PNG of ``side``×``side`` (noise: comfortably over 1200 B)."""
    if color is not None:
        image = Image.new("RGB", (side, side), color)
    else:
        payload = random.Random(seed).randbytes(side * side * 3)
        image = Image.frombytes("RGB", (side, side), payload)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", compress_level=0)
    data = buffer.getvalue()
    if target is not None:
        assert len(data) <= target, (len(data), target)
        data = data + b"\x00" * (target - len(data))
    return data


def _tiny_png(target: int) -> bytes:
    """A small but *valid* PNG padded to exactly ``target`` bytes."""
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), (11, 22, 33)).save(buffer, format="PNG", compress_level=0)
    data = buffer.getvalue()
    assert len(data) <= target, (len(data), target)
    return data + b"\x00" * (target - len(data))


def _docx_bytes(paragraphs: list[str], images: list[bytes]) -> bytes:
    document = DocxDocument()
    for index, text in enumerate(paragraphs):
        document.add_paragraph(text)
        if index < len(images):
            document.add_picture(io.BytesIO(images[index]), width=Inches(0.5))
    for image in images[len(paragraphs) :]:
        document.add_picture(io.BytesIO(image), width=Inches(0.5))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(slides: list[tuple[str, list[bytes]]]) -> bytes:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for text, images in slides:
        slide = presentation.slides.add_slide(blank)
        if text:
            box = slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1))
            box.text_frame.text = text
        for image in images:
            slide.shapes.add_picture(
                io.BytesIO(image), PptxInches(1), PptxInches(2), width=PptxInches(0.5)
            )
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(pages: list[tuple[str, list[bytes]]]) -> bytes:
    document = pymupdf.open()
    for text, images in pages:
        page = document.new_page()
        if text:
            page.insert_text((36, 36), text)
        for index, image in enumerate(images):
            rect = pymupdf.Rect(36, 60 + index * 70, 36 + 64, 60 + index * 70 + 64)
            page.insert_image(rect, stream=image)
    data = document.tobytes()
    document.close()
    return data


def _record(filename: str, data: bytes, *, record_id: str = "doc-1") -> dict[str, Any]:
    return {
        "type": "document",
        "url": "",
        "base64": base64.b64encode(data).decode("ascii"),
        "filename": filename,
        "mime_type": "application/octet-stream",
        "id": record_id,
    }


def _image_records(updated: list[dict]) -> list[dict]:
    return [record for record in updated if record.get("embedded")]


def _emitted_bytes(records: list[dict]) -> int:
    return sum(len(base64.b64decode(record["base64"])) for record in records)


def _document_text(doc_texts: list[str], filename: str) -> str:
    prefix = f"[File: {filename}]\n"
    matching = [text for text in doc_texts if text.startswith(prefix)]
    assert matching, doc_texts
    return matching[0][len(prefix) :]


def _assert_stats_consistent(doc_texts: list[str], updated: list[dict], filename: str) -> str:
    """``extracted_chars`` and the attachment note must match reality."""
    text = _document_text(doc_texts, filename)
    records = _image_records(updated)
    document_records = [
        record
        for record in updated
        if record.get("filename") == filename and record.get("type") != "image"
    ]
    assert len(document_records) == 1
    assert document_records[0]["extracted_chars"] == len(text)
    assert document_records[0]["extracted_text"] == text
    assert document_records[0]["base64"] == ""
    stem = filename.rsplit(".", 1)[0]
    own_records = [item for item in records if item["filename"].startswith(f"{stem}-图")]
    if own_records:
        assert (
            f"本文件内嵌 {len(own_records)} 张图片" in text
        ), "the note must state the real emitted attachment count for this document"
    return text


# ---------------------------------------------------------------------------
# 1. 三格式贯通：正文 + 有效图 → 标记 / record / 来源 / 顺序
# ---------------------------------------------------------------------------
def test_docx_text_and_image_records_match_source_and_markers(p1_runtime: Path) -> None:
    images = [_png_bytes(seed=1), _png_bytes(seed=2)]
    data = _docx_bytes(["第一章 引言", "第二章 方法"], images)
    record = _record("report.docx", data, record_id="doc-a")

    doc_texts, updated = extractor.extract_documents_from_records([record])

    text = _assert_stats_consistent(doc_texts, updated, "report.docx")
    assert "第一章 引言" in text and "第二章 方法" in text
    markers = find_markers(text)
    assert [index for index, _ in markers] == [1, 2], "markers keep document order"
    assert [name for _, name in markers] == ["image-01.png", "image-02.png"]

    records = _image_records(updated)
    assert [item["filename"] for item in records] == ["report-图1.png", "report-图2.png"]
    assert [item["id"] for item in records] == ["doc-a-e01", "doc-a-e02"]
    assert [item["mime_type"] for item in records] == [_MIME, _MIME]
    assert [len(base64.b64decode(item["base64"])) for item in records] == [
        len(images[0]),
        len(images[1]),
    ]
    assert _emitted_bytes(records) == sum(len(image) for image in images)
    # The images travel as ordinary image attachments of this document.
    assert all(item["type"] == "image" for item in records)


def test_pptx_text_and_image_records_match_source_and_markers(p1_runtime: Path) -> None:
    images = [_png_bytes(seed=3)]
    data = _pptx_bytes([("第一页要点", images), ("第二页要点", [])])
    record = _record("deck.pptx", data, record_id="doc-b")

    doc_texts, updated = extractor.extract_documents_from_records([record])

    text = _assert_stats_consistent(doc_texts, updated, "deck.pptx")
    assert "--- Slide 1 ---" in text and "第一页要点" in text
    assert find_markers(text) == [(1, "image-01.png")]
    records = _image_records(updated)
    assert [item["filename"] for item in records] == ["deck-图1.png"]
    assert [item["id"] for item in records] == ["doc-b-e01"]
    assert len(base64.b64decode(records[0]["base64"])) == len(images[0])


def test_pdf_markers_land_on_the_page_that_holds_the_image(p1_runtime: Path) -> None:
    page1_image = _png_bytes(seed=4)
    page2_image = _png_bytes(seed=5)
    data = _pdf_bytes([("page one text", [page1_image]), ("page two text", [page2_image])])
    record = _record("paper.pdf", data, record_id="doc-c")

    doc_texts, updated = extractor.extract_documents_from_records([record])

    text = _assert_stats_consistent(doc_texts, updated, "paper.pdf")
    page_one, _, page_two = text.partition("--- Page 2 ---")
    assert find_markers(page_one) == [(1, "image-01.png")]
    assert find_markers(page_two) == [(2, "image-02.png")]

    records = _image_records(updated)
    assert [item["filename"] for item in records] == ["paper-图1.png", "paper-图2.png"]
    assert [item["id"] for item in records] == ["doc-c-e01", "doc-c-e02"]
    assert _emitted_bytes(records) > 0
    decoded = [base64.b64decode(item["base64"]) for item in records]
    # PyMuPDF re-encodes the embedded stream, so the emitted bytes are asserted
    # to be real, decodable images (and distinct), not byte-identical copies.
    for payload in decoded:
        with Image.open(io.BytesIO(payload)) as image:
            image.verify()
    assert decoded[0] != decoded[1]


def test_text_only_document_is_unchanged(p1_runtime: Path) -> None:
    data = _docx_bytes(["第一章 只有文字内容"], [])
    doc_texts, updated = extractor.extract_documents_from_records([_record("plain.docx", data)])

    text = _assert_stats_consistent(doc_texts, updated, "plain.docx")
    assert text == "第一章 只有文字内容"
    assert _image_records(updated) == []
    assert "[图片" not in text and "内嵌" not in text


# ---------------------------------------------------------------------------
# 2. 预算：每文档 / 每轮 / 字节边界 / 跨文档共享
# ---------------------------------------------------------------------------
def test_document_image_cap_allows_twelve_and_reports_thirteen(p1_runtime: Path) -> None:
    images = [_png_bytes(seed=index) for index in range(MAX_IMAGES_PER_DOC + 1)]
    data = _docx_bytes(["正文"], images)
    doc_texts, updated = extractor.extract_documents_from_records([_record("many.docx", data)])

    text = _assert_stats_consistent(doc_texts, updated, "many.docx")
    records = _image_records(updated)
    assert len(records) == MAX_IMAGES_PER_DOC == 12
    assert [item["filename"] for item in records] == [f"many-图{n}.png" for n in range(1, 13)]
    assert len(find_markers(text)) == 12
    # The 13th candidate really is in the document (13 media parts) and really
    # was left behind: the collection holds 12 and counts one budget skip.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 13
    collection = extract_docx_rich(data).collection
    assert len(collection.images) == 12
    assert collection.skipped_budget == 1
    assert "1 张图片超出数量/大小上限未提取" in text


def test_turn_image_cap_is_shared_across_documents(p1_runtime: Path) -> None:
    first = _docx_bytes(["A"], [_png_bytes(seed=index) for index in range(10)])
    second = _docx_bytes(["B"], [_png_bytes(seed=100 + index) for index in range(10)])
    records_in = [
        _record("first.docx", first, record_id="d1"),
        _record("second.docx", second, record_id="d2"),
    ]

    doc_texts, updated = extractor.extract_documents_from_records(records_in)

    images = _image_records(updated)
    assert len(images) == extractor.MAX_EMBEDDED_IMAGES_PER_TURN == 16
    assert [item["id"] for item in images][:10] == [f"d1-e{n:02d}" for n in range(1, 11)]
    assert [item["id"] for item in images][10:] == [f"d2-e{n:02d}" for n in range(1, 7)]

    second_text = _assert_stats_consistent(doc_texts, updated, "second.docx")
    assert "本文件还有 4 张图片因数量/大小上限未随消息提供" in second_text
    assert len(find_markers(second_text)) == 10, "all ten markers stay in the extracted text"
    first_text = _assert_stats_consistent(doc_texts, updated, "first.docx")
    assert len(find_markers(first_text)) == 10
    assert "未随消息提供" not in first_text


def test_per_image_byte_boundary_is_inclusive(p1_runtime: Path) -> None:
    """1199 rejected, 1200 kept; 4 MiB kept, 4 MiB + 1 rejected."""
    under = _tiny_png(MIN_IMAGE_BYTES - 1)
    exact_min = _tiny_png(MIN_IMAGE_BYTES)
    exact_max = _png_bytes(MAX_IMAGE_BYTES, seed=9)
    over_max = _png_bytes(MAX_IMAGE_BYTES + 1, seed=10)
    assert [len(item) for item in (under, exact_min, exact_max, over_max)] == [
        MIN_IMAGE_BYTES - 1,
        MIN_IMAGE_BYTES,
        MAX_IMAGE_BYTES,
        MAX_IMAGE_BYTES + 1,
    ]

    kept = _docx_bytes(["边界"], [under, exact_min])
    kept_texts, kept_updated = extractor.extract_documents_from_records(
        [_record("kept.docx", kept)]
    )
    kept_records = _image_records(kept_updated)
    assert [len(base64.b64decode(item["base64"])) for item in kept_records] == [
        MIN_IMAGE_BYTES
    ]
    kept_text = _assert_stats_consistent(kept_texts, kept_updated, "kept.docx")
    assert "1 张图片超出数量/大小上限未提取" in kept_text

    big = _docx_bytes(["大图"], [exact_max, over_max])
    big_texts, big_updated = extractor.extract_documents_from_records([_record("big.docx", big)])
    big_records = _image_records(big_updated)
    assert [len(base64.b64decode(item["base64"])) for item in big_records] == [MAX_IMAGE_BYTES]
    big_text = _assert_stats_consistent(big_texts, big_updated, "big.docx")
    assert "1 张图片超出数量/大小上限未提取" in big_text


def test_document_total_byte_cap_and_turn_byte_cap(p1_runtime: Path) -> None:
    """Per-document 12 MiB and per-turn 12 MiB are both enforced on real bytes."""
    per_image = MAX_IMAGE_BYTES
    doc_a = _docx_bytes(["A"], [_png_bytes(per_image, seed=20 + index) for index in range(4)])
    doc_b = _docx_bytes(["B"], [_png_bytes(per_image, seed=30)])
    records_in = [_record("a.docx", doc_a, record_id="a"), _record("b.docx", doc_b, record_id="b")]

    doc_texts, updated = extractor.extract_documents_from_records(records_in)

    images = _image_records(updated)
    assert len(images) == 3, "the turn cap stops after 12 MiB"
    assert _emitted_bytes(images) == MAX_TOTAL_IMAGE_BYTES_PER_DOC == 12 * 1024 * 1024
    a_text = _assert_stats_consistent(doc_texts, updated, "a.docx")
    assert "1 张图片超出数量/大小上限未提取" in a_text  # 4th is over the *document* cap
    b_text = _assert_stats_consistent(doc_texts, updated, "b.docx")
    assert "本文件还有 1 张图片因数量/大小上限未随消息提供" in b_text  # blocked by the *turn* cap
    assert "1 张图片超出数量/大小上限未提取" not in b_text


def test_reading_budget_is_untouched_and_separate() -> None:
    budget = reading_image_budget()
    assert budget.max_images == 120
    assert budget.max_total_bytes == 60 * 1024 * 1024
    assert budget.max_images_per_page == 4
    default = ImageBudget()
    assert default.max_images == MAX_IMAGES_PER_DOC == 12
    assert default.max_total_bytes == MAX_TOTAL_IMAGE_BYTES_PER_DOC
    assert default.max_image_bytes == MAX_IMAGE_BYTES
    assert default.min_image_bytes == MIN_IMAGE_BYTES == 1200
    assert default.max_images_per_page == 0


def test_reading_budget_per_page_cap_still_applies(p1_runtime: Path) -> None:
    images = [_png_bytes(seed=40 + index) for index in range(6)]
    data = _pdf_bytes([("page", images)])

    default_result = extract_pdf_images(data)
    assert len(default_result.collection.images) == 6
    assert default_result.collection.skipped_page_budget == 0

    reading_result = extract_pdf_images(data, budget=reading_image_budget())
    assert len(reading_result.collection.images) == 4
    assert reading_result.collection.skipped_page_budget == 2


# ---------------------------------------------------------------------------
# 3. 重复资产 / 非支持与矢量格式
# ---------------------------------------------------------------------------
def test_repeated_asset_is_stored_once_and_marked_at_every_position(p1_runtime: Path) -> None:
    image = _png_bytes(seed=50)
    data = _docx_bytes(["第一处", "第二处"], [image, image])
    record = _record("repeat.docx", data, record_id="rep")

    doc_texts, updated = extractor.extract_documents_from_records([record])

    text = _assert_stats_consistent(doc_texts, updated, "repeat.docx")
    markers = find_markers(text)
    assert markers == [(1, "image-01.png"), (1, "image-01.png")], "both positions stay marked"
    records = _image_records(updated)
    assert len(records) == 1
    assert len(base64.b64decode(records[0]["base64"])) == len(image)
    rich = extract_docx_rich(data)
    assert len(rich.collection.images) == 1


def test_vector_and_unsupported_assets_are_reported_not_emitted(p1_runtime: Path) -> None:
    data = _minimal_docx_with_media(
        [
            ("word/media/pic.emf", b"\x01\x00\x00\x00EMF-PLACEHOLDER" + b"\x00" * 2000),
            ("word/media/scan.tiff", _png_bytes(seed=60)),
        ]
    )
    doc_texts, updated = extractor.extract_documents_from_records([_record("mix.docx", data)])

    text = _assert_stats_consistent(doc_texts, updated, "mix.docx")
    assert _image_records(updated) == []
    assert "1 张矢量图(EMF/WMF)未提取" in text
    assert "1 张不支持的图片格式未提取" in text
    rich = extract_docx_rich(data)
    assert rich.collection.skipped_vector == 1
    assert rich.collection.skipped_other == 1


def test_external_relationship_is_never_fetched(p1_runtime: Path, tmp_path: Path) -> None:
    outside = tmp_path / "outside-secret.png"
    outside.write_bytes(_png_bytes(seed=70))
    data = _minimal_docx_with_relationships(
        [
            (
                "rId9",
                outside.as_uri(),
                "External",
            )
        ],
        blip_rids=["rId9"],
    )
    before = sorted(path.name for path in tmp_path.iterdir())

    doc_texts, updated = extractor.extract_documents_from_records([_record("ext.docx", data)])

    assert _image_records(updated) == []
    assert sorted(path.name for path in tmp_path.iterdir()) == before
    _assert_stats_consistent(doc_texts, updated, "ext.docx")


def test_relationship_escaping_the_package_is_rejected(p1_runtime: Path) -> None:
    """A ``../../`` target must not reach a member outside the OOXML package.

    The archive really does contain a member stored as ``../evil.png``: without a
    package-root constraint the relationship walk reads it and offers it as a
    figure of ``trav.docx``.
    """
    payload = _png_bytes(seed=80)
    data = _minimal_docx_with_relationships(
        [("rId7", "../../evil.png", None)],
        blip_rids=["rId7"],
        extra_members={"../evil.png": payload},
    )
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "../evil.png" in archive.namelist(), "the fixture must contain the member"

    doc_texts, updated = extractor.extract_documents_from_records([_record("trav.docx", data)])

    assert _image_records(updated) == []
    text = _assert_stats_consistent(doc_texts, updated, "trav.docx")
    assert "evil.png" not in text and "image-01" not in text
    assert extract_docx_rich(data).collection.images == ()


# ---------------------------------------------------------------------------
# 4. 坏 ZIP / 坏 XML：不失败消息、不产出图片
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "filename,payload",
    [
        ("broken.docx", b"this is definitely not a zip archive"),
        ("broken.pptx", b"PK\x03\x04 truncated"),
        ("broken.pdf", b"%PDF-1.4\nnot really a pdf"),
    ],
)
def test_corrupt_container_does_not_fail_the_message(
    p1_runtime: Path, filename: str, payload: bytes
) -> None:
    doc_texts, updated = extractor.extract_documents_from_records([_record(filename, payload)])

    assert _image_records(updated) == []
    record = [item for item in updated if item.get("filename") == filename][0]
    assert record["extracted_chars"] == 0
    assert record["base64"] == ""
    assert doc_texts and "could not be read" in doc_texts[0]


def test_corrupt_ooxml_xml_does_not_fail_the_message(p1_runtime: Path) -> None:
    data = _minimal_zip(
        {
            "[Content_Types].xml": b"<Types/>",
            "word/document.xml": b"<w:document><broken",
            "word/_rels/document.xml.rels": b"<Relationships>",
        }
    )
    doc_texts, updated = extractor.extract_documents_from_records([_record("badxml.docx", data)])

    assert _image_records(updated) == []
    record = [item for item in updated if item.get("filename") == "badxml.docx"][0]
    assert record["extracted_chars"] == 0
    assert doc_texts and "could not be read" in doc_texts[0]


def test_extraction_never_writes_outside_the_process(p1_runtime: Path, tmp_path: Path) -> None:
    """The whole pipeline runs in memory: no host paths are created."""
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    data = _docx_bytes(["正文"], [_png_bytes(seed=90)])
    fixture = workspace / "inside.docx"
    fixture.write_bytes(data)
    snapshot = {path.relative_to(workspace).as_posix() for path in workspace.rglob("*")}

    doc_texts, updated = extractor.extract_documents_from_records(
        [_record("inside.docx", data)]
    )

    assert _image_records(updated)
    assert {path.relative_to(workspace).as_posix() for path in workspace.rglob("*")} == snapshot
    assert _assert_stats_consistent(doc_texts, updated, "inside.docx")


def test_network_guard_is_active_during_extraction(p1_runtime: Path) -> None:
    """Positive control: the guard blocks an external connect *and* extraction works."""
    import socket

    probe = socket.socket()
    try:
        with pytest.raises(helpers.ExternalNetworkBlocked):
            probe.connect(("93.184.216.34", 80))
    finally:
        probe.close()

    data = _docx_bytes(["正文"], [_png_bytes(seed=91)])
    doc_texts, updated = extractor.extract_documents_from_records([_record("net.docx", data)])
    assert _image_records(updated)
    assert _assert_stats_consistent(doc_texts, updated, "net.docx")


# ---------------------------------------------------------------------------
# minimal OOXML builders for the adversarial cases
# ---------------------------------------------------------------------------
def _minimal_zip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def _docx_document_xml(blip_rids: list[str], text: str = "正文") -> bytes:
    blips = "".join(
        f'<w:p><w:r><w:drawing><a:blip r:embed="{rid}"/></w:drawing></w:r></w:p>'
        for rid in blip_rids
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p>{blips}</w:body></w:document>"
    ).encode("utf-8")


def _minimal_docx_with_relationships(
    relationships: list[tuple[str, str, str | None]],
    *,
    blip_rids: list[str],
    extra_members: dict[str, bytes] | None = None,
) -> bytes:
    entries = "".join(
        f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/'
        f'officeDocument/2006/relationships/image" Target="{target}"'
        + (f' TargetMode="{mode}"' if mode else "")
        + "/>"
        for rid, target, mode in relationships
    )
    members = {
        "[Content_Types].xml": b"<Types/>",
        "word/document.xml": _docx_document_xml(blip_rids),
        "word/_rels/document.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            f'relationships">{entries}</Relationships>'
        ).encode("utf-8"),
    }
    members.update(extra_members or {})
    return _minimal_zip(members)


def _minimal_docx_with_media(media: list[tuple[str, bytes]]) -> bytes:
    relationships = [
        (f"rId{index}", f"media/{name.rsplit('/', 1)[-1]}", None)
        for index, (name, _payload) in enumerate(media, start=1)
    ]
    return _minimal_docx_with_relationships(
        relationships,
        blip_rids=[f"rId{index}" for index in range(1, len(media) + 1)],
        extra_members={name: payload for name, payload in media},
    )


# ---------------------------------------------------------------------------
# 5. S06 返修 R6：PDF 像素门不再静默丢图
# ---------------------------------------------------------------------------
def _tiny_pixel_pdf(*, pages: int = 1, side: int = 16) -> bytes:
    """A PDF that really contains ``pages`` placements of one tiny raster."""
    stream = io.BytesIO()
    Image.new("RGB", (side, side), (7, 8, 9)).save(stream, format="PNG", compress_level=0)
    raw = stream.getvalue()
    payload = raw + b"\x00" * (1600 - len(raw))
    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page()
        page.insert_image(pymupdf.Rect(20, 20, 100, 100), stream=payload)
    data = document.tobytes()
    document.close()
    return data


def test_pdf_small_pixel_image_is_counted_as_a_skip(p1_runtime: Path) -> None:
    """The 64x64 heuristic stays, but a rejected xref is a *named* skip."""
    result = extract_pdf_images(_tiny_pixel_pdf())

    assert result.collection.images == ()
    assert result.page_map == ()
    counted = (
        result.collection.skipped_budget
        + result.collection.skipped_other
        + result.collection.skipped_vector
    )
    assert counted == 1, "an observed size-rejected figure silently disappeared"
    assert result.collection.skipped_budget == 1


def test_pdf_reused_small_xref_is_counted_once(p1_runtime: Path) -> None:
    """The same rejected xref on two pages is one observed image, one skip."""
    result = extract_pdf_images(_tiny_pixel_pdf(pages=2))

    assert result.collection.images == ()
    assert result.collection.skipped_budget == 1, "a reused xref must not be double counted"


def test_pdf_byte_small_and_reused_images_keep_their_counts(p1_runtime: Path) -> None:
    """Byte-small figures are skips; a reused valid figure keeps both pages."""
    tiny_bytes = io.BytesIO()
    Image.new("RGB", (64, 64), (9, 9, 9)).save(tiny_bytes, format="PNG")
    assert len(tiny_bytes.getvalue()) < MIN_IMAGE_BYTES

    document = pymupdf.open()
    page_one = document.new_page()
    page_one.insert_image(pymupdf.Rect(10, 10, 200, 200), stream=tiny_bytes.getvalue())
    page_two = document.new_page()
    page_two.insert_image(pymupdf.Rect(10, 10, 200, 200), stream=_png_bytes(seed=11))
    page_two.insert_image(pymupdf.Rect(220, 10, 420, 200), stream=_png_bytes(seed=12))
    data = document.tobytes()
    document.close()

    result = extract_pdf_images(data)
    assert len(result.collection.images) == 2
    assert result.collection.skipped_budget == 1, "the byte-small figure is one skip"
    assert sorted(page for page, _indices in result.page_map) == [2]
    assert result.collection.summary_note().count("超出数量/大小上限") == 1


def test_pdf_reused_valid_xref_keeps_every_page_position(p1_runtime: Path) -> None:
    shared = _png_bytes(seed=13)
    document = pymupdf.open()
    for _ in range(3):
        page = document.new_page()
        page.insert_image(pymupdf.Rect(10, 10, 200, 200), stream=shared)
    data = document.tobytes()
    document.close()

    result = extract_pdf_images(data)
    assert len(result.collection.images) == 1
    assert [page for page, _indices in result.page_map] == [1, 2, 3]
    assert result.collection.skipped_budget == 0
