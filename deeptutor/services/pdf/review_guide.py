"""Review-pack generation (M6) — aggregate the subject's exam points,
wrong-question book and knowledge map into a printable review guide.

Pipeline: aggregate data -> LLM writes Markdown (Unicode math, no LaTeX)
-> renderer emits PDF (Typst) or DOCX (python-docx).
"""

from __future__ import annotations

import logging
from pathlib import Path
import subprocess

from deeptutor.services.llm.config import get_llm_config
from deeptutor.subjects import exam_points as ep_mod
from deeptutor.subjects import knowledge_points, wrong_book
from deeptutor.subjects.generation import complete_subject_text

logger = logging.getLogger(__name__)

GUIDE_PROMPT = """你是一名大学课程助教，为一名学生编写【期末复习指南】。下面给出该科目的：确认过的期末考点、错题集（含知识点与掌握状态）、知识点索引。请编写一份中文复习指南。

## 要求
1. 结构：`# 期末复习指南` 开头，然后是 `## 一、期末考点速览`（逐条列出考点，标注[必考]/[重点]）、`## 二、错题回顾与易错点`（按掌握度从低到高，每条给出：题目要点、易错原因、正确思路）、`## 三、复习建议`（按优先级给出 3-5 条）。
2. 数学公式用 Unicode 文本（如 y′ + P(x)y = Q(x)、e^(−∫P dx)、x²），不要用 LaTeX 命令（不要用 \\frac、\\int、^ 上标命令等）。
3. 只使用 Markdown 的子集：#、##、- 列表、**加粗**。不要用表格、代码块、链接。
4. 内容具体、可执行，直接基于给定的考点与错题，不要泛泛而谈。
"""


def _aggregate(subject_id: str, term: str) -> dict:
    exam_points = [p for p in ep_mod.load(subject_id) if p.get("confirmed")]
    if term:
        exam_points = [p for p in exam_points if p["term"] == term]
    wrongs = [e for e in wrong_book.load(subject_id) if not e.get("archived")]
    points = knowledge_points.aggregate(subject_id, term=term or None)
    return {"exam_points": exam_points, "wrong_entries": wrongs, "knowledge_points": points}


def _format_input(data: dict) -> str:
    lines = ["## 期末考点"]
    for p in data["exam_points"]:
        lines.append(
            f"- [{p['level']}] {p['content']}"
            + (f"（证据：{p['evidence'][:60]}）" if p.get("evidence") else "")
        )
    lines.append("")
    lines.append("## 错题集")
    order = {"repeated": 0, "missed": 1, "basic": 2, "mastered": 3}
    wrongs = sorted(data["wrong_entries"], key=lambda e: order.get(e["status"], 9))
    for e in wrongs:
        lines.append(
            f"- [{e['status']}] 知识点: {e.get('kp') or '未标注'} | 题目: {e['question'][:80]}"
            + (f" | 答案要点: {e['answer'][:60]}" if e.get("answer") else "")
        )
    lines.append("")
    lines.append("## 知识点索引")
    for p in data["knowledge_points"][:15]:
        lines.append(
            f"- {p['name']}（错题 {p['wrong_count']} 道，"
            f"掌握度 {knowledge_points.format_mastery(p['mastery'])}）"
        )
    return "\n".join(lines)


async def generate_guide_markdown(subject_id: str, term: str = "") -> str:
    """LLM-generated review guide (Markdown subset).

    Raises ``LLMReasoningBudgetExhausted`` when the provider stopped at an
    output cap (the reasoning model burned the budget), or
    ``LLMEmptyResponseError`` for any other empty result: a blank guide must
    never be rendered into a PDF/DOCX that looks like a successful generation.
    """
    data = _aggregate(subject_id, term)
    if not data["exam_points"] and not data["wrong_entries"]:
        return "# 期末复习指南\n\n该科目暂无已确认的考点与错题，请先上传材料并确认考点。"
    cfg = get_llm_config()
    return await complete_subject_text(
        prompt=GUIDE_PROMPT + "\n\n## 科目数据\n" + _format_input(data),
        system_prompt="你是课程助教，输出 Markdown 复习指南。",
        model=cfg.model,
        api_key=cfg.api_key,
        base_url=cfg.base_url,
    )


# ---------------------- renderers ----------------------


#: Typst markup specials that must be escaped in plain body text. The report
#: and guide generators are told to write math as Unicode text, so bodies
#: routinely contain ``<``/``>`` (inequalities such as ``0<a<1``). Typst reads
#: ``<…>`` as a *label*, so one unmatched ``<`` aborts the whole compile with
#: "unclosed label" — previously surfacing as an opaque render failure.
_TYPST_ESCAPES = {
    "\\": "\\\\",
    "#": "\\#",
    "$": "\\$",
    "<": "\\<",
    ">": "\\>",
    "@": "\\@",
    "[": "\\[",
    "]": "\\]",
    "`": "\\`",
    "*": "\\*",
    "_": "\\_",
}


def _escape_typst(text: str) -> str:
    """Escape Typst markup specials in a run of plain text."""
    return "".join(_TYPST_ESCAPES.get(ch, ch) for ch in text)


def _md_to_typst(md: str) -> str:
    """Convert the Markdown subset (headings/lists/bold) to Typst markup."""
    out = []
    in_list = False
    for line in md.splitlines():
        stripped = line.rstrip()
        if not stripped.strip():
            if in_list:
                out.append("")
                in_list = False
            continue
        if stripped.startswith("### "):
            if in_list:
                out.append("")
                in_list = False
            out.append(f"== {_escape_typst(stripped[4:].strip())}")
        elif stripped.startswith("## "):
            if in_list:
                out.append("")
                in_list = False
            out.append(f"= {_escape_typst(stripped[3:].strip())}")
        elif stripped.startswith("# "):
            if in_list:
                out.append("")
                in_list = False
            title = _escape_typst(stripped[2:].strip())
            out.append(f'#align(center)[#text(size: 22pt, weight: "bold")[{title}]]')
        elif stripped.startswith("- ") or stripped.startswith("* "):
            if not in_list:
                out.append("")
                in_list = True
            item = _escape_typst(stripped[2:].strip().replace("**", ""))
            out.append(f"- {item}")
        else:
            if in_list:
                out.append("")
                in_list = False
            body = _escape_typst(stripped.replace("**", ""))
            # A body line opening with "=" or "+" would be read as a Typst
            # heading / list marker rather than text.
            if body[:1] in ("=", "+") and body[1:2] == " ":
                body = "\\" + body
            out.append(body)
    return "\n".join(out)


_TYST_DOC = """#set page(paper: \"a4\", margin: 2cm)
#set text(size: 11pt, lang: \"zh\", font: \"Microsoft YaHei\")
#set par(justify: true)

{body}
"""

_FONT_DIR = Path(__file__).resolve().parents[3] / "tools" / "fonts"


def typst_compile(typ_path: Path, out_path: Path, typst_bin: Path | None = None) -> bool:
    """Compile a .typ file to PDF with the project CJK font path.

    Centralizes font-path wiring so every renderer (review guide, report)
    picks up the Windows CJK fonts (Microsoft YaHei) shipped in tools/fonts.
    """
    typst_bin = typst_bin or (Path(__file__).resolve().parents[3] / "tools" / "bin" / "typst")
    if not typst_bin.exists():
        return False
    try:
        proc = subprocess.run(
            [
                str(typst_bin),
                "compile",
                "--font-path",
                str(_FONT_DIR),
                str(typ_path),
                str(out_path),
            ],
            capture_output=True,
            timeout=120,
        )
        if proc.returncode != 0:
            logger.warning("typst compile failed: %s", proc.stderr.decode(errors="replace")[:300])
        return proc.returncode == 0 and out_path.exists()
    except Exception:
        return False


def render_pdf(markdown: str, out_path: Path, typst_bin: Path | None = None) -> bool:
    """Compile Markdown to PDF via Typst. Returns success."""
    src = _TYST_DOC.format(body=_md_to_typst(markdown))
    tmp = out_path.with_suffix(".typ")
    tmp.write_text(src, encoding="utf-8")
    return typst_compile(tmp, out_path, typst_bin)


def render_docx(markdown: str, out_path: Path) -> bool:
    """Render Markdown subset to DOCX via python-docx."""
    from docx import Document

    doc = Document()
    in_list = False
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            in_list = False
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith(("- ", "* ")):
            if not in_list:
                doc.add_paragraph()
                in_list = True
            doc.add_paragraph(stripped[2:].replace("**", ""), style="List Bullet")
        else:
            in_list = False
            doc.add_paragraph(stripped.replace("**", ""))
    doc.save(str(out_path))
    return out_path.exists()


def estimate_tokens(text: str) -> int:
    """Rough token estimate via tiktoken (cl100k_base)."""
    try:
        import tiktoken

        enc = tiktoken.get_encoding("cl100k_base")
        return len(enc.encode(text))
    except Exception:
        return max(1, len(text) // 3)
